"""Render a structured Resume as plain text, DOCX or PDF."""

import io
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import HRFlowable, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from .schemas import Resume


def _join(*parts: str, sep: str = " | ") -> str:
    return sep.join(p for p in parts if p and p.strip())


def to_text(resume: Resume) -> str:
    out: list[str] = []
    if resume.name:
        out.append(resume.name)
    if resume.headline:
        out.append(resume.headline)
    if resume.contact:
        out.append(_join(*resume.contact))

    def section(title: str) -> None:
        out.extend(["", title.upper()])

    if resume.summary:
        section("Summary")
        out.append(resume.summary)
    if resume.skills:
        section("Skills")
        out.append(", ".join(resume.skills))
    if resume.experience:
        section("Experience")
        for job in resume.experience:
            out.append(_join(job.title, job.company, job.location, job.dates))
            out.extend(f"• {b}" for b in job.bullets)
            out.append("")
        out.pop()
    if resume.education:
        section("Education")
        for edu in resume.education:
            out.append(_join(edu.degree, edu.institution, edu.dates))
            if edu.details:
                out.append(edu.details)
    for extra in resume.extra_sections:
        section(extra.heading)
        out.extend(f"• {item}" for item in extra.items)
    return "\n".join(out).strip() + "\n"


def to_docx(resume: Resume) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Pt(50)

    def centered(text: str, size: float, bold: bool = False) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.bold, run.font.size = bold, Pt(size)

    def heading(text: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text.upper())
        run.bold, run.font.size = True, Pt(11.5)
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    def bullet(text: str) -> None:
        p = doc.add_paragraph(text, style="List Bullet")
        p.paragraph_format.space_after = Pt(1)

    if resume.name:
        centered(resume.name, 18, bold=True)
    if resume.headline:
        centered(resume.headline, 11)
    if resume.contact:
        centered(_join(*resume.contact), 9.5)

    if resume.summary:
        heading("Summary")
        doc.add_paragraph(resume.summary)
    if resume.skills:
        heading("Skills")
        doc.add_paragraph(", ".join(resume.skills))
    if resume.experience:
        heading("Experience")
        for job in resume.experience:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1)
            p.add_run(_join(job.title, job.company, sep=" — ")).bold = True
            meta = _join(job.location, job.dates)
            if meta:
                p.add_run(f"   {meta}").italic = True
            for b in job.bullets:
                bullet(b)
    if resume.education:
        heading("Education")
        for edu in resume.education:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.add_run(_join(edu.degree, edu.institution, sep=" — ")).bold = True
            if edu.dates:
                p.add_run(f"   {edu.dates}").italic = True
            if edu.details:
                doc.add_paragraph(edu.details)
    for extra in resume.extra_sections:
        heading(extra.heading)
        for item in extra.items:
            bullet(item)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf(resume: Resume) -> bytes:
    base = getSampleStyleSheet()
    accent = colors.HexColor("#1E3A8A")
    styles = {
        "name": ParagraphStyle("name", parent=base["Title"], fontSize=18, leading=22, spaceAfter=2),
        "center": ParagraphStyle("center", parent=base["Normal"], alignment=TA_CENTER, fontSize=9.5, leading=12),
        "h": ParagraphStyle("h", parent=base["Heading2"], fontSize=11.5, textColor=accent,
                            spaceBefore=10, spaceAfter=3),
        "body": ParagraphStyle("body", parent=base["Normal"], fontSize=10, leading=13.5),
        "role": ParagraphStyle("role", parent=base["Normal"], fontSize=10, leading=13.5, spaceBefore=5),
    }
    # reportlab's Paragraph accepts a mini-markup, so user/LLM text must be escaped.
    e = escape
    story: list = []

    if resume.name:
        story.append(Paragraph(e(resume.name), styles["name"]))
    if resume.headline:
        story.append(Paragraph(e(resume.headline), styles["center"]))
    if resume.contact:
        story.append(Paragraph(e(_join(*resume.contact)), styles["center"]))

    def heading(text: str) -> None:
        story.append(Paragraph(e(text.upper()), styles["h"]))
        story.append(HRFlowable(width="100%", thickness=0.6, color=accent, spaceAfter=4))

    def bullets(items: list[str]) -> None:
        if items:
            story.append(ListFlowable(
                [ListItem(Paragraph(e(i), styles["body"]), leftIndent=12) for i in items],
                bulletType="bullet", start="•", leftIndent=12, bulletFontSize=8,
            ))

    if resume.summary:
        heading("Summary")
        story.append(Paragraph(e(resume.summary), styles["body"]))
    if resume.skills:
        heading("Skills")
        story.append(Paragraph(e(", ".join(resume.skills)), styles["body"]))
    if resume.experience:
        heading("Experience")
        for job in resume.experience:
            line = f"<b>{e(_join(job.title, job.company, sep=' — '))}</b>"
            meta = _join(job.location, job.dates)
            if meta:
                line += f"  <i>{e(meta)}</i>"
            story.append(Paragraph(line, styles["role"]))
            bullets(job.bullets)
    if resume.education:
        heading("Education")
        for edu in resume.education:
            line = f"<b>{e(_join(edu.degree, edu.institution, sep=' — '))}</b>"
            if edu.dates:
                line += f"  <i>{e(edu.dates)}</i>"
            story.append(Paragraph(line, styles["role"]))
            if edu.details:
                story.append(Paragraph(e(edu.details), styles["body"]))
    for extra in resume.extra_sections:
        heading(extra.heading)
        bullets(extra.items)

    if not story:
        story.append(Spacer(1, 1))
    buf = io.BytesIO()
    SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=15 * mm, bottomMargin=15 * mm,
        title=f"{resume.name} - Resume".strip(" -"),
    ).build(story)
    return buf.getvalue()
