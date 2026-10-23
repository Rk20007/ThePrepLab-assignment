"""Extract plain text from uploaded resume files."""

import io
import logging
import re
import zipfile

import pdfplumber
from docx import Document

# pdfminer logs a warning for every minor spec deviation, which real-world resumes have plenty of.
logging.getLogger("pdfminer").setLevel(logging.ERROR)


class ResumeParseError(ValueError):
    """Raised when a file can't be turned into usable resume text."""


SUPPORTED_EXTENSIONS = (".pdf", ".docx")


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        if not data.startswith(b"%PDF"):
            raise ResumeParseError("This file has a .pdf extension but isn't a valid PDF.")
        text = _extract_pdf(data)
    elif name.endswith(".docx"):
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise ResumeParseError("This file has a .docx extension but isn't a valid Word document.")
        text = _extract_docx(data)
    else:
        raise ResumeParseError("Unsupported file type. Please upload a PDF or DOCX file.")

    text = _normalize(text)
    if len(text) < 50:
        raise ResumeParseError(
            "We couldn't find readable text in this file. If it's a scanned image, "
            "please export it as a text-based PDF or DOCX."
        )
    return text


def _extract_pdf(data: bytes) -> str:
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [page.extract_text(x_tolerance=1.5) or "" for page in pdf.pages]
    except Exception as exc:  # pdfminer raises a variety of exception types
        raise ResumeParseError("The PDF appears to be corrupted or password-protected.") from exc
    return "\n".join(pages)


def _extract_docx(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeParseError("The Word document appears to be corrupted.") from exc

    lines = [p.text for p in doc.paragraphs]
    # Many resume templates put content inside tables.
    for table in doc.tables:
        for row in table.rows:
            seen: set[int] = set()
            for cell in row.cells:
                if id(cell._tc) in seen:  # merged cells repeat
                    continue
                seen.add(id(cell._tc))
                lines.extend(p.text for p in cell.paragraphs)
    return "\n".join(lines)


def _normalize(text: str) -> str:
    text = text.replace(" ", " ").replace("", "•")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()
